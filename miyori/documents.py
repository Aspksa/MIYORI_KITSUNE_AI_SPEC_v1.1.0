from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from xml.etree import ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from .config import settings

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".json",
    ".pdf", ".docx", ".xlsx", ".pptx",
}
MAX_FILE_BYTES = 25 * 1024 * 1024
CHUNK_SIZE = 1400
CHUNK_OVERLAP = 180

DRIVE_DIR_NAME = "Документы - Облако - Miyori"
FILES_DIR_NAME = "Файлы"
TRASH_DIR_NAME = "Корзина"

_WINDOWS_RESERVED = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def project_drive_dir(project_id: int) -> Path:
    """Физическое хранилище Miyori Drive внутри рабочей папки проекта."""
    path = settings.data_dir / "workspace" / str(project_id) / DRIVE_DIR_NAME
    (path / FILES_DIR_NAME).mkdir(parents=True, exist_ok=True)
    (path / TRASH_DIR_NAME).mkdir(parents=True, exist_ok=True)
    return path


def project_document_dir(project_id: int) -> Path:
    """Совместимое имя для корня активных файлов проекта."""
    return project_drive_dir(project_id) / FILES_DIR_NAME


def drive_relative_root(project_id: int) -> str:
    return str(project_drive_dir(project_id).relative_to(settings.data_dir)).replace("\\", "/")


def safe_filename(filename: str) -> str:
    name = Path(filename).name.strip() or "document.txt"
    clean = re.sub(r"[^a-zA-Zа-яА-ЯёЁ0-9._ ()\-]+", "_", name).rstrip(" .")
    return clean or "document.txt"


def safe_folder_name(name: str) -> str:
    clean = " ".join(str(name or "").strip().split())
    if not clean:
        raise ValueError("Название папки пустое.")
    if len(clean) > 120:
        raise ValueError("Название папки слишком длинное.")
    if re.search(r'[<>:"/\\|?*]', clean):
        raise ValueError('В названии папки нельзя использовать символы < > : " / \\ | ? *.')
    if clean.endswith((".", " ")):
        raise ValueError("Название папки не может заканчиваться точкой или пробелом.")
    stem = clean.split(".", 1)[0].upper()
    if stem in _WINDOWS_RESERVED:
        raise ValueError("Это имя зарезервировано Windows.")
    return clean


def _data_relative(path: Path) -> str:
    root = settings.data_dir.resolve()
    resolved = path.resolve()
    if resolved != root and root not in resolved.parents:
        raise ValueError("Путь выходит за пределы каталога данных Miyori.")
    return str(resolved.relative_to(root)).replace("\\", "/")


def resolve_data_path(relative_path: str) -> Path:
    root = settings.data_dir.resolve()
    candidate = (root / relative_path).resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError("Путь выходит за пределы каталога данных Miyori.")
    return candidate


def ensure_drive_folder(project_id: int, folder_parts: list[str] | tuple[str, ...] | None = None) -> Path:
    target = project_document_dir(project_id)
    for part in folder_parts or []:
        target = target / safe_folder_name(part)
    target.mkdir(parents=True, exist_ok=True)
    return target


def _unique_file_target(directory: Path, filename: str, digest: str | None = None) -> Path:
    safe = safe_filename(filename)
    target = directory / safe
    if not target.exists():
        return target

    stem = Path(safe).stem
    suffix = Path(safe).suffix
    token = (digest or "")[:8]
    if token:
        candidate = directory / f"{stem} ({token}){suffix}"
        if not candidate.exists():
            return candidate

    index = 2
    while True:
        candidate = directory / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate
        index += 1


def _unique_dir_target(directory: Path, name: str) -> Path:
    safe = safe_folder_name(name)
    target = directory / safe
    if not target.exists():
        return target
    index = 2
    while True:
        candidate = directory / f"{safe} ({index})"
        if not candidate.exists():
            return candidate
        index += 1


def _decode_text(filename: str, data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Текстовый документ должен быть в UTF-8.") from exc

    if Path(filename).suffix.lower() == ".json":
        try:
            obj = json.loads(text)
            text = json.dumps(obj, ensure_ascii=False, indent=2)
        except json.JSONDecodeError as exc:
            raise ValueError("JSON-файл содержит синтаксическую ошибку.") from exc
    return text


def _decode_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Не удалось открыть PDF.") from exc

    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.strip()
        if text:
            pages.append(f"[PDF · страница {index}]\n{text}")

    if not pages:
        raise ValueError(
            "В PDF не найден извлекаемый текст. "
            "Сканированные PDF пока требуют отдельного OCR."
        )
    return "\n\n".join(pages)


def _decode_docx(data: bytes) -> str:
    try:
        document = WordDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Не удалось открыть DOCX.") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[DOCX · таблица {table_index}]\n" + "\n".join(rows))

    if not parts:
        raise ValueError("В DOCX не найден текст.")
    return "\n\n".join(parts)


def _decode_xlsx(data: bytes) -> str:
    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError("Не удалось открыть XLSX.") from exc

    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append(" | ".join(values))
            if rows:
                parts.append(f"[XLSX · лист {sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()

    if not parts:
        raise ValueError("В XLSX не найдено данных.")
    return "\n\n".join(parts)


def _decode_pptx(data: bytes) -> str:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Не удалось открыть PPTX.") from exc

    parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                slide_parts.append(text.strip())

            if getattr(shape, "has_table", False):
                table_rows: list[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
                if table_rows:
                    slide_parts.append("\n".join(table_rows))

        if slide_parts:
            parts.append(
                f"[PPTX · слайд {slide_index}]\n" +
                "\n".join(slide_parts)
            )

    if not parts:
        raise ValueError("В PPTX не найден текст.")
    return "\n\n".join(parts)



@dataclass
class DocumentBlock:
    index: int
    kind: str
    text: str
    locator: str
    title: str | None = None
    level: int = 0
    metadata: dict = field(default_factory=dict)


@dataclass
class StructuredDocument:
    filename: str
    format: str
    title: str
    text: str
    blocks: list[DocumentBlock]
    metadata: dict = field(default_factory=dict)


def _extraction_manifest(
    *,
    status: str = "complete",
    coverage: float = 1.0,
    warnings: list[str] | None = None,
    details: dict | None = None,
) -> dict:
    return {
        "status": status,
        "coverage": max(0.0, min(1.0, float(coverage))),
        "warnings": list(warnings or []),
        "details": dict(details or {}),
    }


def _zip_inventory(data: bytes, prefix: str) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return [
                name for name in archive.namelist()
                if name.startswith(prefix)
            ]
    except (zipfile.BadZipFile, OSError):
        return []


def _zip_xml_texts(data: bytes, names: list[str]) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            available = set(archive.namelist())
            for name in names:
                if name not in available:
                    continue
                try:
                    root = ET.fromstring(archive.read(name))
                except (ET.ParseError, KeyError):
                    continue
                parts = [
                    str(node.text or "").strip()
                    for node in root.iter()
                    if node.tag.rsplit("}", 1)[-1] == "t"
                    and str(node.text or "").strip()
                ]
                text = " ".join(parts).strip()
                if text:
                    result.append((name, text))
    except (zipfile.BadZipFile, OSError):
        pass
    return result


_HEADING_PREFIX_RE = re.compile(
    r"^(глава|раздел|часть|приложение|chapter|section|part|appendix)\b",
    re.IGNORECASE,
)
_NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+){0,5})[.)]?\s+\S")


def _looks_like_heading(line: str) -> tuple[bool, int]:
    clean = " ".join(line.strip().split())
    if not clean or len(clean) > 180:
        return False, 0

    markdown = re.match(r"^(#{1,6})\s+(.+)$", clean)
    if markdown:
        return True, len(markdown.group(1))

    numbered = _NUMBERED_HEADING_RE.match(clean)
    if numbered:
        return True, min(6, numbered.group(1).count(".") + 1)

    prefix = _HEADING_PREFIX_RE.match(clean)
    if prefix:
        word = prefix.group(1).lower()
        return True, 1 if word in {"глава", "часть", "chapter", "part"} else 2

    letters = [char for char in clean if char.isalpha()]
    if (
        4 <= len(letters) <= 90
        and len(clean.split()) <= 14
        and sum(1 for char in letters if char.isupper()) / max(1, len(letters)) >= 0.82
    ):
        return True, 2

    return False, 0


def _append_block(
    blocks: list[DocumentBlock],
    kind: str,
    text: str,
    locator: str,
    *,
    title: str | None = None,
    level: int = 0,
    metadata: dict | None = None,
) -> None:
    clean = re.sub(r"[ \t]+\n", "\n", str(text or "")).strip()
    if not clean:
        return
    blocks.append(
        DocumentBlock(
            index=len(blocks),
            kind=kind,
            text=clean,
            locator=locator,
            title=(title or "").strip() or None,
            level=max(0, int(level)),
            metadata=dict(metadata or {}),
        )
    )


def _blocks_from_lines(
    text: str,
    locator_prefix: str,
    *,
    metadata: dict | None = None,
    markdown: bool = False,
) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    paragraph: list[str] = []
    paragraph_start = 1
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")

    def flush(end_line: int) -> None:
        nonlocal paragraph
        if paragraph:
            _append_block(
                blocks,
                "paragraph",
                "\n".join(paragraph),
                f"{locator_prefix}:lines:{paragraph_start}-{max(paragraph_start, end_line)}",
                metadata=metadata,
            )
            paragraph = []

    for line_no, raw in enumerate(lines, start=1):
        clean = raw.strip()
        if not clean:
            flush(line_no - 1)
            paragraph_start = line_no + 1
            continue

        is_heading, level = _looks_like_heading(clean)
        if markdown and clean.startswith("#"):
            match = re.match(r"^(#{1,6})\s+(.+)$", clean)
            if match:
                clean = match.group(2).strip()
                is_heading, level = True, len(match.group(1))

        if is_heading:
            flush(line_no - 1)
            _append_block(
                blocks,
                "heading",
                clean,
                f"{locator_prefix}:line:{line_no}",
                title=clean,
                level=level,
                metadata=metadata,
            )
            paragraph_start = line_no + 1
            continue

        if not paragraph:
            paragraph_start = line_no
        paragraph.append(clean)

        if sum(len(item) + 1 for item in paragraph) >= 5000:
            flush(line_no)
            paragraph_start = line_no + 1

    flush(len(lines))
    return blocks


def _structured_text(filename: str, data: bytes) -> StructuredDocument:
    text = _decode_text(filename, data)
    suffix = Path(filename).suffix.lower()
    blocks = _blocks_from_lines(
        text,
        "text",
        markdown=suffix in {".md", ".markdown"},
    )
    title = next((item.title for item in blocks if item.title), None) or Path(filename).stem
    canonical = "\n\n".join(item.text for item in blocks).strip()
    return StructuredDocument(
        filename=filename,
        format=suffix.lstrip("."),
        title=title,
        text=canonical,
        blocks=blocks,
        metadata={
            "page_count": 0,
            "table_count": 0,
            "extraction": _extraction_manifest(
                status="complete",
                coverage=1.0,
                details={"source": "plain_text"},
            ),
        },
    )


def _structured_json(filename: str, data: bytes) -> StructuredDocument:
    try:
        raw = data.decode("utf-8-sig")
        value = json.loads(raw)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("JSON-файл содержит синтаксическую ошибку или неверную кодировку.") from exc

    blocks: list[DocumentBlock] = []
    if isinstance(value, dict):
        for key, item in value.items():
            rendered = json.dumps(item, ensure_ascii=False, indent=2)
            _append_block(
                blocks,
                "section",
                f"{key}\n{rendered}",
                f"json:$.{key}",
                title=str(key),
                level=1,
                metadata={"json_path": f"$.{key}"},
            )
    elif isinstance(value, list):
        for index in range(0, len(value), 50):
            part = value[index:index + 50]
            _append_block(
                blocks,
                "table",
                json.dumps(part, ensure_ascii=False, indent=2),
                f"json:$[{index}:{index + len(part)}]",
                title=f"Элементы {index + 1}–{index + len(part)}",
                level=1,
                metadata={"json_start": index, "json_end": index + len(part) - 1},
            )
    else:
        _append_block(
            blocks,
            "paragraph",
            json.dumps(value, ensure_ascii=False, indent=2),
            "json:$",
        )

    canonical = "\n\n".join(item.text for item in blocks).strip()
    return StructuredDocument(
        filename=filename,
        format="json",
        title=Path(filename).stem,
        text=canonical,
        blocks=blocks,
        metadata={
            "page_count": 0,
            "table_count": 0,
            "extraction": _extraction_manifest(
                status="complete",
                coverage=1.0,
                details={"source": "json_tree"},
            ),
        },
    )


def _structured_pdf(filename: str, data: bytes) -> StructuredDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Не удалось открыть PDF.") from exc

    blocks: list[DocumentBlock] = []
    extracted_pages = 0
    missing_text_pages: list[int] = []
    for page_index, page in enumerate(reader.pages, start=1):
        try:
            page_text = (page.extract_text() or "").strip()
        except Exception:
            page_text = ""
        if not page_text:
            missing_text_pages.append(page_index)
            continue
        extracted_pages += 1
        page_blocks = _blocks_from_lines(
            page_text,
            f"pdf:page:{page_index}",
            metadata={"page": page_index},
        )
        if not page_blocks:
            _append_block(
                blocks,
                "page",
                page_text,
                f"pdf:page:{page_index}",
                title=f"Страница {page_index}",
                metadata={"page": page_index},
            )
        else:
            for item in page_blocks:
                item.index = len(blocks)
                blocks.append(item)

    if not blocks:
        raise ValueError(
            "В PDF не найден извлекаемый текст. "
            "Оригинал можно хранить в Drive, но для полного понимания нужен OCR."
        )

    page_count = len(reader.pages)
    page_coverage = extracted_pages / max(1, page_count)
    warnings: list[str] = []
    if missing_text_pages:
        preview = ", ".join(str(item) for item in missing_text_pages[:30])
        suffix = "…" if len(missing_text_pages) > 30 else ""
        warnings.append(
            "На страницах без извлекаемого текста нужен OCR/визуальный анализ: "
            f"{preview}{suffix}"
        )

    title = next((item.title for item in blocks if item.title), None) or Path(filename).stem
    canonical = "\n\n".join(item.text for item in blocks).strip()
    return StructuredDocument(
        filename=filename,
        format="pdf",
        title=title,
        text=canonical,
        blocks=blocks,
        metadata={
            "page_count": page_count,
            "extracted_page_count": extracted_pages,
            "table_count": 0,
            "extraction": _extraction_manifest(
                status="complete" if not missing_text_pages else "partial",
                coverage=page_coverage,
                warnings=warnings,
                details={
                    "page_count": page_count,
                    "text_pages": extracted_pages,
                    "missing_text_pages": missing_text_pages,
                    "visual_semantics_supported": False,
                },
            ),
        },
    )

def _structured_docx(filename: str, data: bytes) -> StructuredDocument:
    try:
        document = WordDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Не удалось открыть DOCX.") from exc

    blocks: list[DocumentBlock] = []
    paragraph_index = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_index += 1
        style_name = str(getattr(paragraph.style, "name", "") or "")
        match = re.search(r"(?:Heading|Заголовок)\s*(\d+)", style_name, re.IGNORECASE)
        if match:
            level = max(1, min(6, int(match.group(1))))
            _append_block(
                blocks,
                "heading",
                text,
                f"docx:paragraph:{paragraph_index}",
                title=text,
                level=level,
                metadata={"style": style_name},
            )
        elif style_name.lower() in {"title", "название"}:
            _append_block(
                blocks,
                "heading",
                text,
                f"docx:paragraph:{paragraph_index}",
                title=text,
                level=1,
                metadata={"style": style_name},
            )
        else:
            _append_block(
                blocks,
                "paragraph",
                text,
                f"docx:paragraph:{paragraph_index}",
                metadata={"style": style_name},
            )

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            _append_block(
                blocks,
                "table",
                "\n".join(rows),
                f"docx:table:{table_index}",
                title=f"Таблица {table_index}",
                metadata={"table": table_index, "rows": len(rows)},
            )

    seen_header_footer: set[str] = set()
    header_footer_count = 0
    for section_index, section in enumerate(document.sections, start=1):
        for kind, container in (("header", section.header), ("footer", section.footer)):
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in container.paragraphs
                if paragraph.text.strip()
            ).strip()
            if not text or text in seen_header_footer:
                continue
            seen_header_footer.add(text)
            header_footer_count += 1
            _append_block(
                blocks,
                kind,
                text,
                f"docx:section:{section_index}:{kind}",
                title=("Колонтитул" if kind == "header" else "Нижний колонтитул"),
                metadata={"section": section_index, "kind": kind},
            )

    auxiliary_xml = _zip_xml_texts(
        data,
        [
            "word/footnotes.xml",
            "word/endnotes.xml",
            "word/comments.xml",
        ],
    )
    for xml_name, text in auxiliary_xml:
        kind = Path(xml_name).stem
        _append_block(
            blocks,
            kind,
            text,
            f"docx:{kind}",
            title=kind,
            metadata={"source_xml": xml_name},
        )

    media_count = len(_zip_inventory(data, "word/media/"))
    chart_count = len(_zip_inventory(data, "word/charts/"))
    warnings: list[str] = []
    if media_count:
        warnings.append(
            f"В DOCX найдено изображений/медиа: {media_count}; их визуальное содержание не распознано."
        )
    if chart_count:
        warnings.append(
            f"В DOCX найдено диаграмм: {chart_count}; их визуальная семантика не интерпретирована."
        )

    if not blocks:
        raise ValueError("В DOCX не найден текст.")

    title = next((item.title for item in blocks if item.kind == "heading"), None) or Path(filename).stem
    canonical = "\n\n".join(item.text for item in blocks).strip()
    status = "text_only" if warnings else "complete"
    return StructuredDocument(
        filename=filename,
        format="docx",
        title=title,
        text=canonical,
        blocks=blocks,
        metadata={
            "page_count": 0,
            "table_count": len(document.tables),
            "paragraph_count": paragraph_index,
            "header_footer_count": header_footer_count,
            "auxiliary_text_parts": len(auxiliary_xml),
            "media_count": media_count,
            "chart_count": chart_count,
            "extraction": _extraction_manifest(
                status=status,
                coverage=1.0,
                warnings=warnings,
                details={
                    "paragraphs": paragraph_index,
                    "tables": len(document.tables),
                    "headers_footers": header_footer_count,
                    "footnotes_endnotes_comments": len(auxiliary_xml),
                    "media_count": media_count,
                    "chart_count": chart_count,
                    "visual_semantics_supported": False,
                },
            ),
        },
    )

def _structured_xlsx(filename: str, data: bytes) -> StructuredDocument:
    try:
        values_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("Не удалось открыть XLSX.") from exc

    blocks: list[DocumentBlock] = []
    total_rows = 0
    formula_count = 0
    nonempty_sheets = 0
    sheet_count = len(values_book.worksheets)
    try:
        for sheet_index, values_sheet in enumerate(values_book.worksheets):
            formulas_sheet = formulas_book.worksheets[sheet_index]
            _append_block(
                blocks,
                "heading",
                values_sheet.title,
                f"xlsx:sheet:{values_sheet.title}",
                title=values_sheet.title,
                level=1,
                metadata={"sheet": values_sheet.title},
            )
            batch: list[str] = []
            batch_start = 1
            current_row = 0
            sheet_has_data = False
            value_rows = values_sheet.iter_rows(values_only=True)
            formula_rows = formulas_sheet.iter_rows(values_only=True)
            for row_index, (value_row, formula_row) in enumerate(
                zip(value_rows, formula_rows),
                start=1,
            ):
                rendered: list[str] = []
                width = max(len(value_row), len(formula_row))
                for col in range(width):
                    value = value_row[col] if col < len(value_row) else None
                    formula = formula_row[col] if col < len(formula_row) else None
                    if isinstance(formula, str) and formula.startswith("="):
                        formula_count += 1
                        if value is None or str(value) == str(formula):
                            rendered.append(f"[formula: {formula}]")
                        else:
                            rendered.append(f"{value} [formula: {formula}]")
                    else:
                        rendered.append("" if value is None else str(value))

                if not any(item.strip() for item in rendered):
                    continue
                sheet_has_data = True
                current_row = row_index
                total_rows += 1
                if not batch:
                    batch_start = row_index
                batch.append(" | ".join(rendered))
                if len(batch) >= 80:
                    _append_block(
                        blocks,
                        "table",
                        "\n".join(batch),
                        f"xlsx:sheet:{values_sheet.title}:rows:{batch_start}-{row_index}",
                        title=f"{values_sheet.title} · строки {batch_start}–{row_index}",
                        level=2,
                        metadata={
                            "sheet": values_sheet.title,
                            "row_start": batch_start,
                            "row_end": row_index,
                        },
                    )
                    batch = []
            if sheet_has_data:
                nonempty_sheets += 1
            if batch:
                _append_block(
                    blocks,
                    "table",
                    "\n".join(batch),
                    f"xlsx:sheet:{values_sheet.title}:rows:{batch_start}-{current_row}",
                    title=f"{values_sheet.title} · строки {batch_start}–{current_row}",
                    level=2,
                    metadata={
                        "sheet": values_sheet.title,
                        "row_start": batch_start,
                        "row_end": current_row,
                    },
                )
    finally:
        values_book.close()
        formulas_book.close()

    if not blocks:
        raise ValueError("В XLSX не найдено данных.")

    media_count = len(_zip_inventory(data, "xl/media/"))
    chart_count = len(_zip_inventory(data, "xl/charts/"))
    warnings: list[str] = []
    if media_count:
        warnings.append(
            f"В XLSX найдено изображений/медиа: {media_count}; их визуальное содержание не распознано."
        )
    if chart_count:
        warnings.append(
            f"В XLSX найдено диаграмм: {chart_count}; их визуальная семантика не интерпретирована."
        )
    sheet_coverage = nonempty_sheets / max(1, sheet_count)
    status = "text_only" if warnings else "complete"

    canonical = "\n\n".join(item.text for item in blocks).strip()
    return StructuredDocument(
        filename=filename,
        format="xlsx",
        title=Path(filename).stem,
        text=canonical,
        blocks=blocks,
        metadata={
            "page_count": 0,
            "table_count": sum(1 for item in blocks if item.kind == "table"),
            "sheet_count": sheet_count,
            "row_count": total_rows,
            "formula_count": formula_count,
            "media_count": media_count,
            "chart_count": chart_count,
            "extraction": _extraction_manifest(
                status=status,
                coverage=sheet_coverage if sheet_count else 1.0,
                warnings=warnings,
                details={
                    "sheets": sheet_count,
                    "nonempty_sheets": nonempty_sheets,
                    "rows": total_rows,
                    "formulas": formula_count,
                    "media_count": media_count,
                    "chart_count": chart_count,
                    "visual_semantics_supported": False,
                },
            ),
        },
    )

def _structured_pptx(filename: str, data: bytes) -> StructuredDocument:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Не удалось открыть PPTX.") from exc

    blocks: list[DocumentBlock] = []
    table_count = 0
    text_slides = 0
    slides_without_text: list[int] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        title = None
        for shape in slide.shapes:
            text = str(getattr(shape, "text", "") or "").strip()
            if text:
                if title is None and len(text) <= 180:
                    title = text.split("\n", 1)[0].strip()
                slide_parts.append(text)

            if getattr(shape, "has_table", False):
                table_count += 1
                rows: list[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    slide_parts.append("\n".join(rows))

        if slide_parts:
            text_slides += 1
            _append_block(
                blocks,
                "slide",
                "\n\n".join(slide_parts),
                f"pptx:slide:{slide_index}",
                title=title or f"Слайд {slide_index}",
                level=1,
                metadata={"slide": slide_index},
            )
        else:
            slides_without_text.append(slide_index)

    note_parts = _zip_xml_texts(
        data,
        sorted(_zip_inventory(data, "ppt/notesSlides/")),
    )
    for note_index, (xml_name, text) in enumerate(note_parts, start=1):
        _append_block(
            blocks,
            "speaker_notes",
            text,
            f"pptx:notes:{note_index}",
            title=f"Заметки докладчика {note_index}",
            metadata={"source_xml": xml_name},
        )

    if not blocks:
        raise ValueError("В PPTX не найден текст.")

    media_count = len(_zip_inventory(data, "ppt/media/"))
    chart_count = len(_zip_inventory(data, "ppt/charts/"))
    warnings: list[str] = []
    if slides_without_text:
        preview = ", ".join(str(item) for item in slides_without_text[:30])
        suffix = "…" if len(slides_without_text) > 30 else ""
        warnings.append(
            "Слайды без извлекаемого текста требуют визуального анализа: "
            f"{preview}{suffix}"
        )
    if media_count:
        warnings.append(
            f"В PPTX найдено изображений/медиа: {media_count}; их визуальное содержание не распознано."
        )
    if chart_count:
        warnings.append(
            f"В PPTX найдено диаграмм: {chart_count}; их визуальная семантика не интерпретирована."
        )

    slide_count = len(presentation.slides)
    slide_coverage = text_slides / max(1, slide_count)
    status = "partial" if slides_without_text else ("text_only" if warnings else "complete")
    canonical = "\n\n".join(item.text for item in blocks).strip()
    return StructuredDocument(
        filename=filename,
        format="pptx",
        title=blocks[0].title or Path(filename).stem,
        text=canonical,
        blocks=blocks,
        metadata={
            "page_count": slide_count,
            "slide_count": slide_count,
            "text_slide_count": text_slides,
            "table_count": table_count,
            "speaker_note_count": len(note_parts),
            "media_count": media_count,
            "chart_count": chart_count,
            "extraction": _extraction_manifest(
                status=status,
                coverage=slide_coverage,
                warnings=warnings,
                details={
                    "slides": slide_count,
                    "slides_with_text": text_slides,
                    "slides_without_text": slides_without_text,
                    "speaker_notes": len(note_parts),
                    "media_count": media_count,
                    "chart_count": chart_count,
                    "visual_semantics_supported": False,
                },
            ),
        },
    )

def extract_structured_document(filename: str, data: bytes) -> StructuredDocument:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Файл слишком большой. Текущий лимит — 25 МБ.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Формат пока не поддерживается. Разрешены: {allowed}. "
            "Оригинал можно хранить в Miyori Drive без индексирования."
        )

    if suffix == ".json":
        result = _structured_json(filename, data)
    elif suffix in {".txt", ".md", ".markdown"}:
        result = _structured_text(filename, data)
    elif suffix == ".pdf":
        result = _structured_pdf(filename, data)
    elif suffix == ".docx":
        result = _structured_docx(filename, data)
    elif suffix == ".xlsx":
        result = _structured_xlsx(filename, data)
    elif suffix == ".pptx":
        result = _structured_pptx(filename, data)
    else:
        raise ValueError("Неподдерживаемый формат документа.")

    result.text = result.text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return result


def decode_document(filename: str, data: bytes) -> str:
    return extract_structured_document(filename, data).text


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def chunk_text(text: str) -> list[str]:
    clean = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not clean:
        return []

    chunks: list[str] = []
    start = 0
    length = len(clean)

    while start < length:
        end = min(length, start + CHUNK_SIZE)
        if end < length:
            split = max(
                clean.rfind("\n\n", start, end),
                clean.rfind(". ", start, end),
                clean.rfind("\n", start, end),
            )
            if split > start + CHUNK_SIZE // 2:
                end = split + 1

        piece = clean[start:end].strip()
        if piece:
            chunks.append(piece)

        if end >= length:
            break
        start = max(end - CHUNK_OVERLAP, start + 1)

    return chunks


def save_original(
    project_id: int,
    filename: str,
    data: bytes,
    digest: str,
    folder_parts: list[str] | tuple[str, ...] | None = None,
) -> str:
    """Сохраняет байт-в-байт оригинал в активном дереве Miyori Drive."""
    directory = ensure_drive_folder(project_id, folder_parts)
    target = _unique_file_target(directory, filename, digest)
    target.write_bytes(data)
    return _data_relative(target)


def move_active_document(
    project_id: int,
    stored_path: str,
    filename: str,
    folder_parts: list[str] | tuple[str, ...] | None = None,
) -> str:
    """Move an active original inside the same project Drive tree."""
    source = resolve_data_path(stored_path)
    if not source.exists() or not source.is_file():
        raise ValueError("Оригинал документа не найден на диске.")

    active_root = project_document_dir(project_id).resolve()
    resolved_source = source.resolve()
    if resolved_source != active_root and active_root not in resolved_source.parents:
        raise ValueError("Документ находится вне активного хранилища проекта.")

    directory = ensure_drive_folder(project_id, folder_parts)
    if source.parent.resolve() == directory.resolve():
        return _data_relative(source)

    target = _unique_file_target(directory, filename)
    shutil.move(str(source), str(target))
    return _data_relative(target)


def move_document_to_trash(
    project_id: int,
    stored_path: str,
    filename: str,
    document_id: int,
) -> str:
    source = resolve_data_path(stored_path)
    if not source.exists() or not source.is_file():
        raise ValueError("Оригинал документа не найден на диске.")

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    trash = project_drive_dir(project_id) / TRASH_DIR_NAME
    target = _unique_file_target(
        trash,
        f"{stamp}_doc-{document_id}_{safe_filename(filename)}",
    )
    shutil.move(str(source), str(target))
    return _data_relative(target)


def restore_document_from_trash(
    project_id: int,
    trash_path: str,
    filename: str,
    folder_parts: list[str] | tuple[str, ...] | None = None,
) -> str:
    source = resolve_data_path(trash_path)
    if not source.exists() or not source.is_file():
        raise ValueError("Оригинал документа не найден в корзине.")

    directory = ensure_drive_folder(project_id, folder_parts)
    target = _unique_file_target(directory, filename)
    shutil.move(str(source), str(target))
    return _data_relative(target)


def move_folder_to_trash(
    project_id: int,
    folder_parts: list[str] | tuple[str, ...],
    folder_id: int,
) -> str:
    parts = [safe_folder_name(part) for part in folder_parts]
    if not parts:
        raise ValueError("Нельзя удалить корень Miyori Drive.")

    source = project_document_dir(project_id).joinpath(*parts)
    if not source.exists() or not source.is_dir():
        raise ValueError("Папка не найдена на диске.")

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    trash = project_drive_dir(project_id) / TRASH_DIR_NAME
    target = _unique_dir_target(trash, f"{stamp}_folder-{folder_id}_{parts[-1]}")
    shutil.move(str(source), str(target))
    return _data_relative(target)


def restore_folder_from_trash(
    project_id: int,
    trash_path: str,
    parent_parts: list[str] | tuple[str, ...],
    folder_name: str,
) -> tuple[str, str]:
    source = resolve_data_path(trash_path)
    if not source.exists() or not source.is_dir():
        raise ValueError("Папка не найдена в корзине.")

    parent = ensure_drive_folder(project_id, parent_parts)
    target = parent / safe_folder_name(folder_name)
    if target.exists():
        raise ValueError(
            "Исходное место папки уже занято. Освободите его перед восстановлением."
        )
    shutil.move(str(source), str(target))
    return _data_relative(target), target.name


def migrate_legacy_document(
    project_id: int,
    stored_path: str,
    filename: str,
    digest: str,
    folder_parts: list[str] | tuple[str, ...] | None = None,
) -> str:
    """Копирует старый оригинал в Miyori Drive без удаления исходной копии."""
    source = resolve_data_path(stored_path)
    active_root = project_document_dir(project_id).resolve()
    try:
        resolved = source.resolve()
        if resolved == active_root or active_root in resolved.parents:
            return _data_relative(resolved)
    except OSError:
        pass

    if not source.exists() or not source.is_file():
        return stored_path

    data = source.read_bytes()
    if sha256_bytes(data) != digest:
        raise ValueError(f"Контрольная сумма оригинала «{filename}» не совпадает.")
    return save_original(project_id, filename, data, digest, folder_parts)
