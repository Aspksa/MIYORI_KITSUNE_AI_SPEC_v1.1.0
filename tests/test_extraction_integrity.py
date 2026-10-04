from __future__ import annotations

import asyncio
import io
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document as WordDocument
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from miyori.config import settings
from miyori.db import (
    add_document,
    create_project,
    init_db,
)
from miyori.document_intelligence import (
    build_local_document_intelligence,
    get_document_intelligence,
    init_document_intelligence_db,
)
from miyori.document_questions import (
    enqueue_exhaustive_document_question,
    init_document_questions_db,
    run_exhaustive_document_question,
)
from miyori.documents import (
    chunk_text,
    extract_structured_document,
    save_original,
    sha256_bytes,
)


class ExtractionIntegrityTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.old_data_dir = settings.data_dir
        self.old_database_path = settings.database_path
        root = Path(self.tmp.name)
        object.__setattr__(settings, "data_dir", root / "data")
        object.__setattr__(
            settings,
            "database_path",
            root / "data" / "miyori.sqlite3",
        )
        settings.data_dir.mkdir(parents=True, exist_ok=True)
        init_db()
        init_document_intelligence_db()
        init_document_questions_db()
        self.project = create_project("Extraction Integrity", kind="work")
        self.project_id = int(self.project["id"])

    def tearDown(self) -> None:
        object.__setattr__(settings, "data_dir", self.old_data_dir)
        object.__setattr__(settings, "database_path", self.old_database_path)
        self.tmp.cleanup()

    def _persist(self, filename: str, data: bytes):
        structured = extract_structured_document(filename, data)
        digest = sha256_bytes(data)
        stored = save_original(
            self.project_id,
            filename,
            data,
            digest,
        )
        document = add_document(
            self.project_id,
            filename,
            stored,
            "application/octet-stream",
            digest,
            len(data),
            chunk_text(structured.text),
        )
        profile = build_local_document_intelligence(
            self.project_id,
            int(document["id"]),
            structured,
        )
        return structured, document, profile

    def test_docx_extracts_headers_footers_and_marks_visual_media(self) -> None:
        document = WordDocument()
        document.add_heading("Регламент", level=1)
        document.add_paragraph("Основной текст документа.")
        section = document.sections[0]
        section.header.paragraphs[0].text = "Служебный верхний колонтитул"
        section.footer.paragraphs[0].text = "Страница служебного документа"

        image = Image.new("RGB", (20, 20), "white")
        image_bytes = io.BytesIO()
        image.save(image_bytes, format="PNG")
        image_bytes.seek(0)
        document.add_picture(image_bytes, width=Inches(0.25))

        stream = io.BytesIO()
        document.save(stream)
        structured, _, profile = self._persist("regulation.docx", stream.getvalue())

        kinds = {block.kind for block in structured.blocks}
        self.assertIn("header", kinds)
        self.assertIn("footer", kinds)
        self.assertIn("Служебный верхний колонтитул", structured.text)
        self.assertEqual(
            structured.metadata["extraction"]["status"],
            "text_only",
        )
        self.assertEqual(structured.metadata["media_count"], 1)
        self.assertEqual(profile["extraction_status"], "text_only")
        self.assertEqual(profile["extraction_coverage"], 1.0)
        self.assertTrue(profile["extraction_warnings"])

    def test_xlsx_preserves_formula_when_cached_value_is_absent(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Расчёт"
        sheet["A1"] = 10
        sheet["A2"] = 20
        sheet["A3"] = "=SUM(A1:A2)"
        stream = io.BytesIO()
        workbook.save(stream)

        structured, _, profile = self._persist("calc.xlsx", stream.getvalue())

        self.assertIn("[formula: =SUM(A1:A2)]", structured.text)
        self.assertEqual(structured.metadata["formula_count"], 1)
        self.assertEqual(profile["extraction_status"], "complete")
        self.assertEqual(profile["extraction_coverage"], 1.0)

    def test_pptx_blank_slide_reduces_extraction_coverage(self) -> None:
        presentation = Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[1])
        first.shapes.title.text = "Слайд с текстом"
        first.placeholders[1].text = "Содержимое доступно parser-у."
        presentation.slides.add_slide(presentation.slide_layouts[6])
        stream = io.BytesIO()
        presentation.save(stream)

        structured, _, profile = self._persist("deck.pptx", stream.getvalue())

        extraction = structured.metadata["extraction"]
        self.assertEqual(extraction["status"], "partial")
        self.assertAlmostEqual(extraction["coverage"], 0.5, places=3)
        self.assertIn(2, extraction["details"]["slides_without_text"])
        self.assertEqual(profile["extraction_status"], "partial")
        self.assertAlmostEqual(profile["extraction_coverage"], 0.5, places=3)

    def test_exhaustive_question_never_overstates_partial_source(self) -> None:
        presentation = Presentation()
        first = presentation.slides.add_slide(presentation.slide_layouts[1])
        first.shapes.title.text = "Условие"
        first.placeholders[1].text = "Извлечённый текст."
        presentation.slides.add_slide(presentation.slide_layouts[6])
        stream = io.BytesIO()
        presentation.save(stream)

        _, document, profile = self._persist("partial-deck.pptx", stream.getvalue())
        self.assertAlmostEqual(profile["extraction_coverage"], 0.5, places=3)

        queued = enqueue_exhaustive_document_question(
            self.project_id,
            int(document["id"]),
            "Есть ли важные условия?",
        )

        async def no_match(title, question, index, total, content):
            return {
                "relevant": False,
                "answer_fragment": "",
                "evidence": [],
                "caveats": [],
            }

        async def synthesis(title, question, results, coverage_ratio):
            return {
                "answer": "В извлечённом тексте дополнительных условий не найдено.",
                "evidence": [],
                "caveats": [],
                "not_found": True,
                "confidence": "high",
            }

        with patch(
            "miyori.document_questions.analyze_document_question_window",
            AsyncMock(side_effect=no_match),
        ), patch(
            "miyori.document_questions.synthesize_exhaustive_document_answer",
            AsyncMock(side_effect=synthesis),
        ):
            result = asyncio.run(
                run_exhaustive_document_question(
                    self.project_id,
                    int(queued["id"]),
                )
            )

        self.assertGreaterEqual(result["coverage_ratio"], 0.995)
        self.assertAlmostEqual(result["extraction_coverage"], 0.5, places=3)
        self.assertAlmostEqual(result["overall_coverage_ratio"], 0.5, places=3)
        self.assertEqual(result["answer"]["confidence"], "medium")
        self.assertTrue(result["answer"]["caveats"])
        self.assertAlmostEqual(
            result["answer"]["overall_coverage_ratio"],
            0.5,
            places=3,
        )

    def test_text_document_has_complete_extraction_manifest(self) -> None:
        structured, _, profile = self._persist(
            "notes.md",
            "# Заголовок\n\nПолный текст.".encode("utf-8"),
        )
        self.assertEqual(structured.metadata["extraction"]["status"], "complete")
        self.assertEqual(structured.metadata["extraction"]["coverage"], 1.0)
        self.assertEqual(profile["extraction_status"], "complete")
        self.assertEqual(profile["extraction_coverage"], 1.0)


class ExtractionIntegrityMigrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.old_data_dir = settings.data_dir
        self.old_database_path = settings.database_path
        root = Path(self.tmp.name)
        object.__setattr__(settings, "data_dir", root / "data")
        object.__setattr__(
            settings,
            "database_path",
            root / "data" / "miyori.sqlite3",
        )
        settings.data_dir.mkdir(parents=True, exist_ok=True)

    def tearDown(self) -> None:
        object.__setattr__(settings, "data_dir", self.old_data_dir)
        object.__setattr__(settings, "database_path", self.old_database_path)
        self.tmp.cleanup()

    def test_0040_document_intelligence_upgrades_in_place(self) -> None:
        with sqlite3.connect(settings.database_path) as conn:
            conn.execute(
                """
                CREATE TABLE projects(
                    id INTEGER PRIMARY KEY,
                    name TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE documents(
                    id INTEGER PRIMARY KEY,
                    project_id INTEGER NOT NULL,
                    filename TEXT NOT NULL,
                    deleted_at TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE tasks(
                    id INTEGER PRIMARY KEY,
                    project_id INTEGER NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE document_questions(
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    project_id INTEGER NOT NULL,
                    document_id INTEGER NOT NULL,
                    task_id INTEGER,
                    source_sha256 TEXT NOT NULL,
                    question TEXT NOT NULL,
                    question_hash TEXT NOT NULL,
                    status TEXT NOT NULL,
                    scanned_chars INTEGER NOT NULL DEFAULT 0,
                    total_chars INTEGER NOT NULL DEFAULT 0,
                    coverage_ratio REAL NOT NULL DEFAULT 0.0,
                    answer_json TEXT NOT NULL DEFAULT '{}',
                    model_id TEXT,
                    engine_version TEXT NOT NULL,
                    last_error TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    finished_at TEXT
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE document_intelligence(
                    document_id INTEGER PRIMARY KEY,
                    project_id INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'indexed',
                    parser_version TEXT NOT NULL,
                    source_sha256 TEXT NOT NULL,
                    title TEXT,
                    document_kind TEXT NOT NULL,
                    language TEXT,
                    char_count INTEGER NOT NULL DEFAULT 0,
                    word_count INTEGER NOT NULL DEFAULT 0,
                    page_count INTEGER NOT NULL DEFAULT 0,
                    section_count INTEGER NOT NULL DEFAULT 0,
                    table_count INTEGER NOT NULL DEFAULT 0,
                    node_count INTEGER NOT NULL DEFAULT 0,
                    analyzed_chars INTEGER NOT NULL DEFAULT 0,
                    coverage_ratio REAL NOT NULL DEFAULT 0.0,
                    summary_short TEXT,
                    summary_long TEXT,
                    outline_json TEXT NOT NULL DEFAULT '[]',
                    keywords_json TEXT NOT NULL DEFAULT '[]',
                    analysis_json TEXT NOT NULL DEFAULT '{}',
                    analysis_model TEXT,
                    last_error TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )

        init_document_intelligence_db()
        init_document_questions_db()

        with sqlite3.connect(settings.database_path) as conn:
            columns = {
                row[1]
                for row in conn.execute(
                    "PRAGMA table_info(document_intelligence)"
                ).fetchall()
            }

        self.assertIn("extraction_status", columns)
        self.assertIn("extraction_coverage", columns)
        self.assertIn("extraction_warnings_json", columns)
        self.assertIn("extraction_details_json", columns)

        with sqlite3.connect(settings.database_path) as conn:
            question_columns = {
                row[1]
                for row in conn.execute(
                    "PRAGMA table_info(document_questions)"
                ).fetchall()
            }

        self.assertIn("extraction_status", question_columns)
        self.assertIn("extraction_coverage", question_columns)
        self.assertIn("overall_coverage_ratio", question_columns)


if __name__ == "__main__":
    unittest.main()
